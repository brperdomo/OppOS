"""The response document as a Word file in Nutrient's submitted-response layout.

Matches the responses Bryan has submitted (ExxonMobil, Tennessee, Cayman): US Letter, margins top 1" / bottom
0.72" / left 3.13" / right 0.89" (the wide left margin is the house style), the Nutrient logo top-left of every
page, a centred footer "Nutrient · <subject> — Page N", Archivo 10 pt body with a 22 pt centred title.

The Markdown this converts is our own (`render_markdown`), so the subset is known: headings (# … ####),
paragraphs, "- " bullets, "1. " numbered items, one pipe table (Appendix A), "---" rules, "> " notes and
inline **bold**, *italic* and `code`.
"""
from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

LOGO = Path(__file__).parent / "assets" / "nutrient-logo.png"
FONT = "Archivo"
BODY_PT, TITLE_PT, H2_PT, H3_PT, H4_PT, FOOTER_PT = 10, 22, 16, 12, 11, 7.5
MARGINS = {"top": 1.0, "bottom": 0.72, "left": 3.13, "right": 0.89}
LOGO_WIDTH_IN, LOGO_OFFSET_IN = 1.45, 0.3  # logo sits in the left margin, 0.3" from the page edge

_INLINE = re.compile(r"(\*\*.+?\*\*|\*(?!\s).+?(?<!\s)\*|`[^`]+`)")


def _font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _inline(paragraph, text: str, size: float = BODY_PT, base_bold: bool = False, base_italic: bool = False) -> None:
    """Add runs for **bold**, *italic* and `code` spans; a trailing two-space line break is honoured."""
    for chunk in text.split("  \n"):
        if paragraph.runs:
            paragraph.add_run().add_break()
        for part in _INLINE.split(chunk):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                _font(paragraph.add_run(part[2:-2]), size, True, base_italic)
            elif part.startswith("`") and part.endswith("`"):
                r = paragraph.add_run(part[1:-1]); _font(r, size, base_bold, base_italic); r.font.name = "Menlo"
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                _font(paragraph.add_run(part[1:-1]), size, base_bold, True)
            else:
                _font(paragraph.add_run(part), size, base_bold, base_italic)


def _page_field(paragraph) -> None:
    """A PAGE field so Word shows the page number."""
    for tag, text in (("begin", None), (None, " PAGE "), ("end", None)):
        run = paragraph.add_run()
        _font(run, FOOTER_PT)
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)


def _bottom_border(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bot.set(qn(k), v)
    b.append(bot); pPr.append(b)


def _setup(doc: Document, subject: str) -> None:
    s = doc.sections[0]
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin, s.bottom_margin = Inches(MARGINS["top"]), Inches(MARGINS["bottom"])
    s.left_margin, s.right_margin = Inches(MARGINS["left"]), Inches(MARGINS["right"])
    s.header_distance, s.footer_distance = Inches(0.35), Inches(0.3)
    # Header: the logo, pulled into the left margin so it sits near the page edge like the house template.
    hp = s.header.paragraphs[0]
    hp.paragraph_format.left_indent = Inches(LOGO_OFFSET_IN - MARGINS["left"])
    if LOGO.exists():
        hp.add_run().add_picture(str(LOGO), width=Inches(LOGO_WIDTH_IN))
    else:
        _font(hp.add_run("Nutrient"), 14, True)
    # Footer: "Nutrient · <subject> — Page N", centred, small.
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(fp.add_run(f"Nutrient · {subject} — Page "), FOOTER_PT)
    _page_field(fp)
    # Base styles
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", TITLE_PT), ("Heading 2", H2_PT), ("Heading 3", H3_PT), ("Heading 4", H4_PT)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
        st.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 24)
        st.paragraph_format.space_after = Pt(6)


def _table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            _inline(p, row[j] if j < len(row) else "", size=8.5, base_bold=(i == 0))
    doc.add_paragraph()


def markdown_to_docx(markdown: str, subject: str, author: str = "", today: date | None = None) -> bytes:
    doc = Document()
    _setup(doc, subject)
    lines = markdown.splitlines()
    i, table_buf = 0, []
    para_buf: list[str] = []

    def flush_para() -> None:
        if para_buf:
            p = doc.add_paragraph()
            _inline(p, "  \n".join(para_buf))
            para_buf.clear()

    def flush_table() -> None:
        nonlocal table_buf
        rows = []
        for ln in table_buf:
            if re.match(r"^\|\s*-{2,}", ln):
                continue  # separator
            rows.append([c.strip() for c in ln.strip().strip("|").split("|")])
        _table(doc, rows)
        table_buf = []

    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            flush_para(); table_buf.append(ln); i += 1; continue
        if table_buf:
            flush_table()
        if not ln.strip():
            flush_para(); i += 1; continue
        if ln.startswith("# "):
            flush_para()
            p = doc.add_paragraph(style="Heading 1"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _inline(p, ln[2:].strip(), TITLE_PT, base_bold=True)
            rule = doc.add_paragraph(); _bottom_border(rule)
        elif ln.startswith("## "):
            flush_para(); p = doc.add_paragraph(style="Heading 2"); _inline(p, ln[3:].strip(), H2_PT, base_bold=True)
        elif ln.startswith("### "):
            flush_para(); p = doc.add_paragraph(style="Heading 3"); _inline(p, ln[4:].strip(), H3_PT, base_bold=True)
        elif ln.startswith("#### "):
            flush_para(); p = doc.add_paragraph(style="Heading 4"); _inline(p, ln[5:].strip(), H4_PT, base_bold=True)
        elif ln.strip() == "---":
            flush_para(); rule = doc.add_paragraph(); _bottom_border(rule)
        elif ln.startswith("- "):
            flush_para(); p = doc.add_paragraph(style="List Bullet"); _inline(p, ln[2:].strip())
        elif re.match(r"^\d+\. ", ln):
            flush_para(); p = doc.add_paragraph(style="List Number"); _inline(p, re.sub(r"^\d+\. ", "", ln).strip())
        elif ln.startswith("> "):
            flush_para(); p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            _inline(p, ln[2:].strip(), base_italic=True)
        elif ln.startswith("*") and ln.endswith("*") and not ln.startswith("**"):
            flush_para(); p = doc.add_paragraph(); _inline(p, ln.strip("*").strip(), base_italic=True)
        else:
            para_buf.append(ln.rstrip())
        i += 1
    flush_para()
    if table_buf:
        flush_table()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
